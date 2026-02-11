import google.generativeai as genai
import PyPDF2
import docx
import json
import random
import os
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib.auth import login, logout as auth_logout, authenticate
from django.contrib.auth.models import User
from django.contrib import messages
from django.conf import settings
from django.http import HttpResponse
from .models import Test, Question, StudentResult
from dotenv import load_dotenv

load_dotenv()

# --- GEMINI API ---
def get_configured_genai():
    keys = getattr(settings, 'GEMINI_KEYS', [])
    if not keys:
        env_key = os.getenv("GEMINI_API_KEYS")
        if env_key: keys = [k.strip() for k in env_key.split(',') if k.strip()]
    if not keys: return None
    genai.configure(api_key=random.choice(keys))
    return genai

def extract_text(file):
    try:
        if file.name.endswith('.pdf'):
            reader = PyPDF2.PdfReader(file)
            return "".join([page.extract_text() for page in reader.pages])
        elif file.name.endswith('.docx'):
            doc = docx.Document(file)
            return "\n".join([para.text for para in doc.paragraphs])
    except: return None
    return ""

# --- 1. DASHBOARD ---
@login_required(login_url='user_login')
def dashboard(request):
    tests = Test.objects.filter(teacher=request.user)
    student_counts = [StudentResult.objects.filter(test=t).count() for t in tests]
    test_titles = [t.title for t in tests]
    
    # Leaderboard (Топ-5 оқушы)
    leaderboard = StudentResult.objects.filter(test__teacher=request.user).order_by('-score')[:5]

    context = {
        'total_tests': tests.count(),
        'total_students': StudentResult.objects.filter(test__teacher=request.user).count(),
        'chart_labels': json.dumps(test_titles),
        'chart_data': json.dumps(student_counts),
        'leaderboard': leaderboard
    }
    return render(request, 'dashboard.html', context)

# --- 2. CREATE (ТЕСТ ЖҮКТЕУ) ---
@login_required(login_url='user_login')
def upload_file(request):
    if request.method == 'POST' and request.FILES.get('document'):
        try:
            title = request.POST.get('title')
            questions_to_show = int(request.POST.get('questions_to_show', 20)) 
            max_students = int(request.POST.get('max_students', 100))
            time = int(request.POST.get('time_limit', 20))
            mode = request.POST.get('mode', 'lite')
            
            text = extract_text(request.FILES['document'])
            if not text: return render(request, 'upload.html', {'error': "Файл оқылмады!"})

            ai = get_configured_genai()
            if not ai: return render(request, 'upload.html', {'error': "API Key қате!"})

            model = ai.GenerativeModel('gemini-flash-latest')
            
            prompt = f"""
            Task: Extract multiple choice questions from the text.
            Format: JSON Array of objects: [{{"question": "Text", "options": ["A", "B", "C", "D"], "correct": 0}}]
            Rules: "correct" index (0=A, 1=B, 2=C, 3=D). If options are missing, add placeholders.
            Text: {text[:15000]}
            """
            
            response = model.generate_content(prompt)
            json_text = response.text.replace("```json", "").replace("```", "").strip()
            
            try:
                data = json.loads(json_text)
            except:
                start = json_text.find('[')
                end = json_text.rfind(']') + 1
                data = json.loads(json_text[start:end])
            
            new_test = Test.objects.create(
                teacher=request.user, 
                title=title, 
                time_limit=time,
                questions_to_show=questions_to_show,
                max_students=max_students,
                mode=mode
            )
            
            for q in data:
                opts = q.get('options', [])
                while len(opts) < 4: opts.append("-")
                Question.objects.create(
                    test=new_test, text=q['question'], 
                    option1=opts[0], option2=opts[1], option3=opts[2], option4=opts[3], 
                    correct_option=q['correct'] + 1
                )
            
            return redirect('history')
        except Exception as e: 
            return render(request, 'upload.html', {'error': f"Қате: {e}"})
            
    return render(request, 'upload.html')

# --- 3. HISTORY & TOOLS ---
@login_required(login_url='user_login')
def history(request):
    my_tests = Test.objects.filter(teacher=request.user).order_by('-id')
    results = StudentResult.objects.filter(test__teacher=request.user).order_by('-date_taken')
    return render(request, 'history.html', {'my_tests': my_tests, 'results': results})

@login_required(login_url='user_login')
def delete_test(request, test_id):
    test = get_object_or_404(Test, id=test_id, teacher=request.user)
    test.delete()
    return redirect('history')

# --- 4. AUTH (ТІРКЕЛУ ЖӘНЕ КІРУ) ---

def register(request):
    if request.method == 'POST':
        u = request.POST.get('username')
        p = request.POST.get('password')
        pc = request.POST.get('password_confirm')
        
        if p != pc: 
            messages.error(request, "Құпиясөздер сәйкес емес!")
            return redirect('register')
        
        if User.objects.filter(username=u).exists(): 
            messages.error(request, "Бұл логин бос емес!")
            return redirect('register')
        
        # МАҢЫЗДЫ ТҮЗЕТУ: password=p деп нақты көрсету керек!
        user = User.objects.create_user(username=u, password=p)
        
        login(request, user)
        return redirect('dashboard')
    return render(request, 'register.html')

def user_login(request):
    if request.method == 'POST':
        u = request.POST.get('username')
        p = request.POST.get('password')
        user = authenticate(request, username=u, password=p)
        if user: 
            login(request, user)
            return redirect('dashboard')
        else: 
            messages.error(request, "Логин немесе құпиясөз қате!")
    return render(request, 'login.html')

def user_logout(request):
    auth_logout(request)
    return redirect('user_login')

@login_required(login_url='user_login')
def profile(request): 
    return render(request, 'profile.html', {'user': request.user})

# --- 5. TEST TAKING (ТАПСЫРУ) ---
def take_test(request, test_id):
    test = get_object_or_404(Test, id=test_id)
    
    current_students = StudentResult.objects.filter(test=test).count()
    if current_students >= test.max_students:
        return HttpResponse("<h1>Өкінішке орай, орын қалмады! (Лимит толды)</h1>")

    if request.method == 'POST':
        student_name = request.POST.get('student_name')
        score = 0
        user_answers = {}
        question_ids = request.POST.getlist('q_ids')
        
        for q_id in question_ids:
            try:
                question = Question.objects.get(id=q_id)
                selected = request.POST.get(f'question_{q_id}')
                if selected:
                    user_answers[str(q_id)] = int(selected)
                    if int(selected) == question.correct_option: score += 1
                else:
                    user_answers[str(q_id)] = None
            except: continue
        
        result = StudentResult.objects.create(
            test=test, student_name=student_name, score=score, 
            total_questions=len(question_ids), student_answers=user_answers
        )
        return redirect('test_result', result_id=result.id)
    
    all_questions = list(test.questions.all())
    if not all_questions: return HttpResponse("Сұрақтар жоқ.")
    
    random.shuffle(all_questions)
    limit = min(test.questions_to_show, len(all_questions))
    return render(request, 'test.html', {'test': test, 'questions': all_questions[:limit]})

def test_result(request, result_id):
    result = get_object_or_404(StudentResult, id=result_id)
    analysis = []
    for q_id, selected_opt in result.student_answers.items():
        try:
            question = Question.objects.get(id=int(q_id))
            analysis.append({
                'question': question, 'selected': selected_opt,
                'is_correct': (selected_opt == question.correct_option),
                'correct_opt': question.correct_option
            })
        except: continue
    return render(request, 'result.html', {'result': result, 'analysis': analysis})